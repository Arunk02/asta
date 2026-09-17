"""Hands, layer one — Astra-class P7.

His ask from the first message of this plan: "even i ask to create a data in
exceel it has to do". A brain decides the content; code writes the file and
then reads back what it wrote. A file that does not check out is never sent.
"""

from __future__ import annotations

import asyncio

import pytest

from app import files, store

ROWS = [["Environment", "Topic", "Partitions"],
        ["dev", "booking.events", 3],
        ["uat", "booking.events", 6],
        ["prod", "booking.events", 12]]


@pytest.fixture(autouse=True)
def _own_folder(tmp_path, monkeypatch):
    monkeypatch.setenv("ASTA_FILES_DIR", str(tmp_path / "out"))
    yield


@pytest.mark.parametrize("kind", ["xlsx", "csv", "md", "docx", "pptx"])
def test_a_table_survives_the_round_trip_in_every_kind(kind):
    made = files.make(kind, "topics per environment", rows=ROWS,
                      title="Topics per environment")
    assert made.path.endswith(f".{kind}") and made.rows == 3
    back = files.read(made.path)["rows"]
    assert [str(c) for c in back[0]] == ROWS[0]
    assert [str(c) for c in back[3]] == [str(c) for c in ROWS[3]]


def test_a_document_carries_its_prose_and_its_table():
    made = files.make("docx", "uat note", rows=ROWS, title="UAT refresh",
                      text="The uat refresh is done.\n\nProd is scheduled for next week.")
    data = files.read(made.path)
    assert "uat refresh is done" in data["text"] and len(data["rows"]) == 4


def test_a_file_that_does_not_say_what_was_asked_is_never_delivered(monkeypatch):
    """The check is the point: a wrong spreadsheet he trusts is worse than none."""
    def half_a_sheet(path, rows, title, text):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(rows[0])            # the header, and nothing else
        wb.save(path)

    monkeypatch.setattr(files, "_xlsx", half_a_sheet)
    with pytest.raises(RuntimeError, match="rows, expected"):
        files.make("xlsx", "topics", rows=ROWS)


def test_nothing_to_write_and_too_much_to_write_are_both_refused():
    with pytest.raises(ValueError, match="nothing to write"):
        files.make("xlsx", "empty")
    with pytest.raises(ValueError, match="more than"):
        files.make("csv", "huge", rows=[["x"]] * (files.MAX_ROWS + 1))
    with pytest.raises(ValueError, match="kind must be"):
        files.make("keynote", "later", rows=ROWS)
    with pytest.raises(ValueError, match="more than a pptx should carry"):
        files.make("pptx", "a whole database", rows=[["x"]] * 400)


def test_the_name_cannot_escape_the_folder():
    made = files.make("csv", "../../etc/passwd", rows=ROWS)
    assert files.folder() == __import__("pathlib").Path(made.path).parent


def test_a_shared_sheet_becomes_rows_to_answer_from():
    made = files.make("xlsx", "shared", rows=ROWS)
    line = files.describe(made.path)
    assert "3 row(s)" in line and "Environment" in line
    assert "no file at" in files.describe("/nope/nothing.xlsx")


def test_it_reaches_his_phone_and_says_so_when_it_cannot(monkeypatch):
    from app import notify
    pushed: list[str] = []
    sent: list[str] = []

    async def wa_document(path, caption=""):
        sent.append(path)
        return bool(sent) and len(sent) == 1        # the first goes, the second does not

    async def push(text, level="info", **kw):
        pushed.append(text)
        return {}

    monkeypatch.setattr(notify, "wa_document", wa_document)
    monkeypatch.setattr(notify, "notify", push)
    made = files.make("xlsx", "topics", rows=ROWS, title="Topics")
    assert asyncio.run(files.deliver(made))["sent"] is True
    # The document IS the message: it arrives with its caption. A second "📎"
    # push repeated it — live on 17 Sep every file reached him twice, the copy
    # spent his daily budget, and four of them turned up again in the digest.
    assert pushed == [], "a delivered document must not be announced a second time"
    assert any("topics.xlsx" in n["text"] for n in store.list_notifications(10)), \
        "…but the bell still records it"
    assert asyncio.run(files.deliver(made))["sent"] is False
    assert "on your Mac" in pushed[-1]
    kinds = [o["outcome"] for o in store.recent_outcomes(10) if o["kind"] == "file"]
    assert kinds == ["written", "delivered"]


def test_the_capability_writes_checks_and_reports(monkeypatch):
    from app import agent, notify

    async def wa_document(path, caption=""):
        return True

    async def push(text, level="info", **kw):
        return {}

    monkeypatch.setattr(notify, "wa_document", wa_document)
    monkeypatch.setattr(notify, "notify", push)
    out = asyncio.run(agent.make_file("xlsx", "topics", ROWS, "Topics per environment"))
    assert out.startswith("Sent to his phone") and "3 row(s) × 3 column(s)" in out
    deck = asyncio.run(agent.make_file("pptx", "slides", ROWS, "Topics per environment"))
    assert deck.startswith("Sent to his phone") and deck.endswith(".pptx)")
    bad = asyncio.run(agent.make_file("keynote", "slides", ROWS))
    assert bad.startswith("Not created")


# --- a deck and a PDF: the same contract, two formats that read back differently


def test_a_pdf_carries_every_cell_and_the_typography_he_writes_with():
    """His text is full of em dashes; the built-in PDF fonts cannot draw one.
    The words still have to arrive — spelled the long way, never dropped."""
    made = files.make("pdf", "topics", rows=ROWS, title="Topics per environment — Q3",
                      text="Where we stand — uat is the one that bites.")
    page = files._flat(files.read(made.path)["text"])
    for cell in ROWS[0] + ["dev", "uat", "prod"]:
        assert cell in page
    assert "Topics per environment - Q3" in page
    assert "uat is the one that bites" in page
    assert files.describe(made.path).startswith(f"{made.path.split('/')[-1]}: 1 page(s)")


def test_a_deck_puts_the_table_on_slides_and_the_prose_in_bullets():
    rows = [["Environment", "Topic"]] + [[f"env{i}", "booking.events"] for i in range(20)]
    made = files.make("pptx", "review", rows=rows, title="Weekly review",
                      text="First point.\n\nSecond point.\n\nThird point.")
    from pptx import Presentation
    deck = Presentation(made.path)
    tables = [sh for sl in deck.slides for sh in sl.shapes if getattr(sh, "has_table", False)]
    assert len(tables) == 2                      # 20 rows do not fit on one slide
    assert sum(len(t.table.rows) - 1 for t in tables) == 20
    text = files.read(made.path)["text"]
    assert "Weekly review" in text and "Second point." in text


@pytest.mark.parametrize("kind", ["pptx", "pdf"])
def test_a_deck_or_a_pdf_that_lost_a_row_is_never_delivered(kind, monkeypatch):
    """Same guard as the spreadsheet, proved for each format on its own terms:
    a PDF is checked by the words on the page, a deck by the cells in its table."""
    original = getattr(files, f"_{kind}")

    def missing_a_row(path, rows, title, text):
        original(path, [rows[0]] + rows[2:], title, text)      # 'dev' quietly dropped

    monkeypatch.setattr(files, f"_{kind}", missing_a_row)
    with pytest.raises(RuntimeError, match="rows, expected|missing the row"):
        files.make(kind, "topics", rows=ROWS, title="Topics")


def test_a_pdf_that_lost_its_prose_is_caught_too(monkeypatch):
    original = files._pdf

    def table_only(path, rows, title, text):
        original(path, rows, title, "")

    monkeypatch.setattr(files, "_pdf", table_only)
    with pytest.raises(RuntimeError, match="missing the text"):
        files.make("pdf", "topics", rows=ROWS, text="Where we stand — uat bites.")
