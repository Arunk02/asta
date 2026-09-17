"""Hands, layer one: a file is WRITTEN by code, and checked before it is sent.

His ask, from the first message of this plan: "even i ask to create a data in
exceel it has to do". A brain deciding the content is right; a brain typing into
a spreadsheet is not — it is slow, unverifiable, and wrong in ways nobody sees
until he opens the file. So the model produces rows, this writes the file, and
then it OPENS WHAT IT WROTE and checks it says what was asked. A file that does
not survive that check is never delivered: a wrong spreadsheet he trusts is
worse than no spreadsheet.

Reads are the other half: a sheet he shares becomes rows Asta can answer from,
rather than a path it can only talk about.
"""

from __future__ import annotations

import csv
import os
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

KINDS = ("xlsx", "csv", "md", "docx", "pptx", "pdf")
#: What a delivered file may weigh, so a runaway generation cannot fill his phone.
MAX_ROWS = 20000
#: A deck and a page are read by eye, not scrolled: past these, the honest answer
#: is "you want a spreadsheet", said before the file is written rather than after.
KIND_MAX_ROWS = {"pptx": 300, "pdf": 2000}
#: Rows per slide and per table slide, so a table is never wider than a room.
ROWS_PER_SLIDE = 12
BULLETS_PER_SLIDE = 6


def folder() -> Path:
    out = Path(os.environ.get("ASTA_FILES_DIR", "~/Asta files")).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    return out


def _safe(name: str, kind: str) -> Path:
    stem = re.sub(r"[^\w .-]+", "", (name or "asta").strip()) or "asta"
    return folder() / f"{stem[:60]}.{kind}"


@dataclass
class Made:
    path: str
    kind: str
    rows: int
    summary: str

    def line(self) -> str:
        return f"{Path(self.path).name} — {self.summary} ({self.path})"


def make(kind: str, name: str, rows: list[list] | None = None,
         title: str = "", text: str = "") -> Made:
    """Write one file and return what was written. Raises if it does not check out.

    rows   a table, first row the header (xlsx, csv, and a table in docx/md)
    text   prose, for docx and md
    """
    kind = (kind or "").lower().lstrip(".")
    if kind not in KINDS:
        raise ValueError(f"kind must be one of {', '.join(KINDS)} — got {kind!r}")
    rows = [list(r) for r in (rows or [])]
    cap = KIND_MAX_ROWS.get(kind, MAX_ROWS)
    if len(rows) > cap:
        raise ValueError(
            f"{len(rows)} rows is more than a {kind} should carry ({cap}) — "
            "ask for xlsx if he wants all of it, or narrow what goes in the deck")
    if not rows and not text.strip():
        raise ValueError("nothing to write — give rows, text, or both")
    path = _safe(name or title or "asta", kind)
    writer = {"xlsx": _xlsx, "csv": _csv, "md": _md, "docx": _docx,
              "pptx": _pptx, "pdf": _pdf}[kind]
    writer(path, rows, title, text)
    made = Made(str(path), kind, max(0, len(rows) - 1 if rows else 0),
                _summary(kind, rows, title, text))
    _check(made, rows, text)
    return made


def _summary(kind: str, rows: list[list], title: str, text: str) -> str:
    if rows:
        body = f"{max(0, len(rows) - 1)} row(s) × {len(rows[0])} column(s)"
    else:
        body = f"{len(text.split())} words"
    return f"{title or kind.upper()}: {body}"


#: A PDF written with the built-in fonts can only carry latin-1, and he writes
#: em dashes and middots constantly — fpdf raises on the first one. Rather than
#: embed a font file or lose the line, spell those characters the long way. The
#: check below compares against the SAME transliteration, so it still proves the
#: content arrived; it just knows an em dash arrives as a hyphen.
_SUBS = {
    "\u2014": "-", "\u2013": "-", "\u2015": "-", "\u00b7": "-", "\u2022": "-",
    "\u2026": "...", "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u2018": "'", "\u2019": "'", "\u2192": "->", "\u2190": "<-",
    "\u2265": ">=", "\u2264": "<=", "\u20ac": "EUR", "\u00a0": " ",
}


def _latin1(text: str) -> str:
    """The same words, in characters a built-in PDF font can actually draw."""
    out = []
    for ch in str(text):
        ch = _SUBS.get(ch, ch)
        try:
            ch.encode("latin-1")
        except UnicodeEncodeError:
            plain = "".join(c for c in unicodedata.normalize("NFKD", ch)
                            if ord(c) < 128)
            ch = plain                       # an emoji normalises to nothing: drop it
        out.append(ch)
    return "".join(out)


def _flat(text: str) -> str:
    """Whitespace collapsed — a PDF wraps lines wherever the page ends, so the
    only fair way to ask "is this sentence in there" is to ignore the wrapping."""
    return " ".join(str(text).split())


def _xlsx(path: Path, rows: list[list], title: str, text: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet1")[:31]
    for row in rows:
        ws.append(["" if c is None else c for c in row])
    if rows:
        for i, cell in enumerate(ws[1], start=1):
            cell.font = Font(bold=True)
            width = max((len(str(r[i - 1])) for r in rows if len(r) >= i), default=10)
            ws.column_dimensions[cell.column_letter].width = min(60, max(10, width + 2))
    if text.strip():
        ws.append([])
        for line in text.strip().splitlines():
            ws.append([line])
    wb.save(path)


def _csv(path: Path, rows: list[list], title: str, text: str) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)


def _md(path: Path, rows: list[list], title: str, text: str) -> None:
    out = [f"# {title}"] if title else []
    if text.strip():
        out += [text.strip(), ""]
    if rows:
        head, body = rows[0], rows[1:]
        out.append("| " + " | ".join(str(c) for c in head) + " |")
        out.append("|" + "---|" * len(head))
        for r in body:
            out.append("| " + " | ".join("" if c is None else str(c) for c in r) + " |")
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def _docx(path: Path, rows: list[list], title: str, text: str) -> None:
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for para in (text or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                table.cell(i, j).text = "" if cell is None else str(cell)
    doc.save(path)


def _pptx(path: Path, rows: list[list], title: str, text: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    paras = [p.strip() for p in (text or "").split("\n\n") if p.strip()]

    opening = deck.slides.add_slide(deck.slide_layouts[0])
    opening.shapes.title.text = title or path.stem.replace("-", " ").title()
    if len(opening.placeholders) > 1:
        opening.placeholders[1].text = paras[0][:250] if paras else ""

    for i in range(0 if not paras else 1, len(paras), BULLETS_PER_SLIDE):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title or "Notes"
        body = slide.placeholders[1].text_frame
        for j, para in enumerate(paras[i:i + BULLETS_PER_SLIDE]):
            line = body.paragraphs[0] if j == 0 else body.add_paragraph()
            line.text = para[:300]

    if rows:
        head, body_rows = rows[0], rows[1:] or []
        chunks = ([body_rows[i:i + ROWS_PER_SLIDE]
                   for i in range(0, len(body_rows), ROWS_PER_SLIDE)] or [[]])
        for chunk in chunks:
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = title or "Data"
            shape = slide.shapes.add_table(
                len(chunk) + 1, len(head), Inches(0.5), Inches(1.7),
                Inches(9), Inches(0.4 * (len(chunk) + 1)))
            table = shape.table
            for j, cell in enumerate(head):
                table.cell(0, j).text = "" if cell is None else str(cell)
            for i, row in enumerate(chunk, start=1):
                for j in range(len(head)):
                    value = row[j] if j < len(row) else ""
                    table.cell(i, j).text = "" if value is None else str(value)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(12)
    deck.save(path)


def _pdf(path: Path, rows: list[list], title: str, text: str) -> None:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    if title:
        pdf.set_font("Helvetica", "B", 16)
        pdf.multi_cell(0, 9, _latin1(title))
        pdf.ln(2)
    if text.strip():
        pdf.set_font("Helvetica", size=11)
        for para in text.strip().split("\n\n"):
            if para.strip():
                pdf.multi_cell(0, 6, _latin1(para.strip()))
                pdf.ln(2)
    if rows:
        pdf.set_font("Helvetica", size=10)
        with pdf.table(line_height=6) as table:
            for row in rows:
                line = table.row()
                for cell in row:
                    line.cell(_latin1("" if cell is None else str(cell)))
    pdf.output(str(path))


def _check(made: Made, rows: list[list], text: str) -> None:
    """Open what was just written and confirm it says what was asked.

    The step that makes this worth trusting. A spreadsheet that silently lost
    its last column, or a document written to a folder that was not writable,
    is exactly the kind of thing he would find hours later in front of someone."""
    path = Path(made.path)
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"{path} was not written")
    back = read(str(path))
    if made.kind == "pdf":
        # A PDF gives back a page of words, not a grid: the cells are all there,
        # but which line they sit on is the page's business. So look for the
        # words — transliterated the same way they were written.
        got = _flat(back.get("text") or "")
        missing = [str(c) for c in (rows[0] if rows else [])
                   if _flat(_latin1(str(c))) not in got]
        if missing:
            raise RuntimeError(
                f"{path.name} is missing header cell(s): {', '.join(missing[:3])}")
        for row in (rows[1:] if len(rows) > 1 else []):
            head = _flat(_latin1(str(row[0]))) if row else ""
            if head and head not in got:
                raise RuntimeError(f"{path.name} is missing the row for {row[0]!r}")
        if text.strip() and _flat(_latin1(text.strip()))[:40] not in got:
            raise RuntimeError(f"{path.name} is missing the text it was given")
        return
    if rows:
        want, got_rows = len(rows), len(back.get("rows") or [])
        if got_rows < want:
            raise RuntimeError(f"{path.name} has {got_rows} rows, expected {want}")
        first = [str(c) for c in rows[0]]
        if [str(c) for c in (back["rows"][0] if back.get("rows") else [])][:len(first)] != first:
            raise RuntimeError(f"{path.name} does not start with the header it was given")
    if text.strip() and text.strip().split("\n")[0][:40] not in (back.get("text") or ""):
        raise RuntimeError(f"{path.name} is missing the text it was given")


def read(path: str, limit: int = 500) -> dict:
    """A file he shared, as data: {rows, text}. Empty for anything unreadable."""
    p = Path(path).expanduser()
    if not p.is_file():
        return {"rows": [], "text": "", "error": f"no file at {p}"}
    kind = p.suffix.lower().lstrip(".")
    try:
        if kind == "xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(p, data_only=True).active
            rows = [[("" if c is None else c) for c in r]
                    for r in ws.iter_rows(max_row=limit, values_only=True)]
            rows = [r for r in rows if any(str(c).strip() for c in r)]
            return {"rows": rows, "text": ""}
        if kind == "csv":
            with p.open(newline="", encoding="utf-8", errors="replace") as fh:
                return {"rows": [r for r in list(csv.reader(fh))[:limit] if r], "text": ""}
        if kind == "pptx":
            from pptx import Presentation
            deck = Presentation(str(p))
            said, rows = [], []
            for slide in deck.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        said.append(shape.text_frame.text.strip())
                    if getattr(shape, "has_table", False):
                        rows += [[c.text for c in r.cells] for r in shape.table.rows]
            return {"rows": rows[:limit], "text": "\n".join(said)}
        if kind == "pdf":
            from pypdf import PdfReader
            pages = PdfReader(str(p)).pages
            text = "\n".join((page.extract_text() or "") for page in pages)
            # No rows: a PDF keeps words, not cells. `describe` says so rather
            # than inventing a grid out of where the lines happened to break.
            return {"rows": [], "text": text[:20000], "pages": len(pages)}
        if kind == "docx":
            from docx import Document
            doc = Document(str(p))
            text = "\n".join(x.text for x in doc.paragraphs)
            rows = [[c.text for c in r.cells] for t in doc.tables for r in t.rows][:limit]
            return {"rows": rows, "text": text}
        text = p.read_text(encoding="utf-8", errors="replace")[:20000]
        # A markdown table is data too — and the check below reads back what it
        # just wrote, so "rows" has to mean the same thing for every kind.
        rows = []
        for line in text.splitlines():
            line = line.strip()
            if line.startswith("|") and not re.fullmatch(r"\|[\s|:-]+\|?", line):
                rows.append([c.strip() for c in line.strip("|").split("|")])
        return {"rows": rows[:limit], "text": text}
    except Exception as exc:                                    # noqa: BLE001
        return {"rows": [], "text": "", "error": f"could not read {p.name}: {exc}"}


async def deliver(made: Made, caption: str = "") -> dict:
    """Put the file where he is: his phone if the bridge will carry it, else a
    line telling him exactly where it is on the Mac. Never silently nowhere."""
    from . import notify, store
    line = caption or made.line()
    sent = await notify.wa_document(made.path, line)
    store.record_outcome("file", "delivered" if sent else "written",
                         subject=Path(made.path).name, detail=made.summary[:200])
    if sent:
        # The document IS the message — it arrives with this line as its caption.
        # Pushing "📎 …" as well sent every file twice; the copy spent his daily
        # budget and, once that ran out, reappeared in the digest (17 Sep). The
        # bell still records it, which is all the second message ever added.
        store.add_notification(f"📎 {line}", "file")
        return {"sent": True, "path": made.path}
    await notify.notify(f"📄 {line}\n\nIt is on your Mac — WhatsApp would not take it.",
                        "file", urgency="direct", considered=True)
    return {"sent": False, "path": made.path}


def describe(path: str) -> str:
    """One line about a file he shared, for a brain to reason with."""
    data = read(path)
    if data.get("error"):
        return data["error"]
    if data["rows"]:
        head = ", ".join(str(c) for c in data["rows"][0][:8])
        return f"{Path(path).name}: {len(data['rows']) - 1} row(s), columns: {head}"
    words = len((data.get("text") or "").split())
    if data.get("pages"):
        return f"{Path(path).name}: {data['pages']} page(s), {words} words"
    return f"{Path(path).name}: {words} words"
